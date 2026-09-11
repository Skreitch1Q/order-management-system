"""Работа с базой данных на основе Excel.

Данные хранятся в книге Excel (файл .xlsx) с отдельными листами:
products, customers, orders, order_items. Модуль отвечает за сохранение
и загрузку данных между запусками приложения, а также за импорт и
экспорт в форматы CSV и JSON.
"""

from __future__ import annotations

import csv
import json
import os
from typing import List, Optional

from openpyxl import Workbook, load_workbook

from models import Customer, Order, Product

# Заголовки колонок для каждого листа книги Excel.
_HEADERS: dict = {
    "products": ["product_id", "name", "price", "category"],
    "customers": ["customer_id", "name", "email", "phone", "city"],
    "orders": ["order_id", "customer_id", "order_date", "status"],
    "order_items": ["item_id", "order_id", "product_id", "quantity"],
}

# Название колонки-идентификатора для каждого листа.
_ID_COLUMNS: dict = {
    "products": "product_id",
    "customers": "customer_id",
    "orders": "order_id",
    "order_items": "item_id",
}


class Database:
    """Управление базой данных в файле Excel.

    Parameters
    ----------
    db_path : str, optional
        Путь к файлу Excel. По умолчанию "shop.xlsx" в текущей директории.
    """

    def __init__(self, db_path: str = "shop.xlsx") -> None:
        self._db_path = db_path
        self._ensure_workbook()

    # ------------------------------------------------------------------ низкий уровень

    def _ensure_workbook(self) -> None:
        """Открыть или создать книгу Excel со всеми необходимыми листами."""
        if os.path.exists(self._db_path):
            self._workbook = load_workbook(self._db_path)
        else:
            self._workbook = Workbook()
            self._workbook.remove(self._workbook.active)
        for sheet, headers in _HEADERS.items():
            ws = self._sheet(sheet, create=True)
            if ws.max_row == 0:
                ws.append(headers)
        self._save()

    def _sheet(self, name: str, create: bool = True):
        """Получить лист книги Excel по имени.

        Parameters
        ----------
        name : str
            Имя листа.
        create : bool, optional
            Создать лист, если он отсутствует (по умолчанию True).

        Returns
        -------
        openpyxl.worksheet.worksheet.Worksheet
            Лист книги.
        """
        if name in self._workbook.sheetnames:
            return self._workbook[name]
        if create:
            return self._workbook.create_sheet(name)
        raise ValueError(f"Лист {name} не найден")

    def _save(self) -> None:
        """Сохранить книгу Excel на диск."""
        self._workbook.save(self._db_path)

    def _read_all(self, name: str) -> List[dict]:
        """Прочитать все строки листа в список словарей.

        Parameters
        ----------
        name : str
            Имя листа.

        Returns
        -------
        list[dict]
            Список строк, каждая строка — словарь "колонка: значение".
        """
        ws = self._sheet(name)
        headers = _HEADERS[name]
        rows = []
        for values in ws.iter_rows(min_row=2, values_only=True):
            if values is None or all(v is None for v in values):
                continue
            rows.append(dict(zip(headers, values)))
        return rows

    def _rewrite(self, name: str, rows: List[dict]) -> None:
        """Полностью перезаписать лист новыми строками.

        Parameters
        ----------
        name : str
            Имя листа.
        rows : list[dict]
            Список строк для записи.
        """
        ws = self._sheet(name)
        ws.delete_rows(1, ws.max_row)
        headers = _HEADERS[name]
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
        self._save()

    def _next_id(self, name: str) -> int:
        """Вычислить следующий свободный идентификатор.

        Parameters
        ----------
        name : str
            Имя листа.

        Returns
        -------
        int
            Следующий идентификатор.
        """
        id_column = _ID_COLUMNS[name]
        ids = [
            int(r[id_column])
            for r in self._read_all(name)
            if r.get(id_column) is not None
        ]
        return (max(ids) + 1) if ids else 1

    # --- Товары ---

    def add_product(self, product: Product) -> int:
        """Добавить товар в базу.

        Parameters
        ----------
        product : Product
            Товар для добавления.

        Returns
        -------
        int
            Идентификатор добавленного товара.
        """
        rows = self._read_all("products")
        if product.product_id is None or any(
            r["product_id"] == product.product_id for r in rows
        ):
            product.product_id = self._next_id("products")
        else:
            rows = [r for r in rows if r["product_id"] != product.product_id]
        rows.append(
            {
                "product_id": product.product_id,
                "name": product.name,
                "price": product.price,
                "category": product.category,
            }
        )
        self._rewrite("products", rows)
        return product.product_id

    def update_product(self, product: Product) -> None:
        """Обновить данные товара в базе.

        Parameters
        ----------
        product : Product
            Товар с обновлёнными данными.
        """
        rows = self._read_all("products")
        for row in rows:
            if row["product_id"] == product.product_id:
                row["name"] = product.name
                row["price"] = product.price
                row["category"] = product.category
        self._rewrite("products", rows)

    def delete_product(self, product_id: int) -> None:
        """Удалить товар по идентификатору.

        Parameters
        ----------
        product_id : int
            Идентификатор товара.
        """
        rows = [r for r in self._read_all("products") if r["product_id"] != product_id]
        self._rewrite("products", rows)

    def get_product(self, product_id: int) -> Optional[Product]:
        """Получить товар по идентификатору.

        Parameters
        ----------
        product_id : int
            Идентификатор товара.

        Returns
        -------
        Product | None
            Товар или None, если товар не найден.
        """
        for row in self._read_all("products"):
            if row["product_id"] == product_id:
                return Product(
                    row["name"],
                    row["price"],
                    row["category"],
                    product_id=row["product_id"],
                )
        return None

    def get_all_products(self) -> List[Product]:
        """Получить список всех товаров.

        Returns
        -------
        list[Product]
            Список товаров, отсортированный по названию.
        """
        rows = sorted(
            self._read_all("products"), key=lambda r: str(r.get("name", ""))
        )
        return [
            Product(r["name"], r["price"], r["category"], product_id=r["product_id"])
            for r in rows
        ]

    # --- Клиенты ---

    def add_customer(self, customer: Customer) -> int:
        """Добавить клиента в базу.

        Parameters
        ----------
        customer : Customer
            Клиент для добавления.

        Returns
        -------
        int
            Идентификатор добавленного клиента.
        """
        rows = self._read_all("customers")
        if customer.customer_id is None or any(
            r["customer_id"] == customer.customer_id for r in rows
        ):
            customer.customer_id = self._next_id("customers")
        else:
            rows = [r for r in rows if r["customer_id"] != customer.customer_id]
        rows.append(
            {
                "customer_id": customer.customer_id,
                "name": customer.name,
                "email": customer.email,
                "phone": customer.phone,
                "city": customer.city,
            }
        )
        self._rewrite("customers", rows)
        return customer.customer_id

    def update_customer(self, customer: Customer) -> None:
        """Обновить данные клиента в базе.

        Parameters
        ----------
        customer : Customer
            Клиент с обновлёнными данными.
        """
        rows = self._read_all("customers")
        for row in rows:
            if row["customer_id"] == customer.customer_id:
                row["name"] = customer.name
                row["email"] = customer.email
                row["phone"] = customer.phone
                row["city"] = customer.city
        self._rewrite("customers", rows)

    def delete_customer(self, customer_id: int) -> None:
        """Удалить клиента по идентификатору.

        Parameters
        ----------
        customer_id : int
            Идентификатор клиента.
        """
        rows = [
            r for r in self._read_all("customers") if r["customer_id"] != customer_id
        ]
        self._rewrite("customers", rows)

    def get_customer(self, customer_id: int) -> Optional[Customer]:
        """Получить клиента по идентификатору.

        Parameters
        ----------
        customer_id : int
            Идентификатор клиента.

        Returns
        -------
        Customer | None
            Клиент или None, если клиент не найден.
        """
        for row in self._read_all("customers"):
            if row["customer_id"] == customer_id:
                return Customer(
                    row["name"],
                    row["email"],
                    row["phone"],
                    row["city"],
                    customer_id=row["customer_id"],
                )
        return None

    def get_all_customers(self) -> List[Customer]:
        """Получить список всех клиентов.

        Returns
        -------
        list[Customer]
            Список клиентов, отсортированный по имени.
        """
        rows = sorted(
            self._read_all("customers"), key=lambda r: str(r.get("name", ""))
        )
        return [
            Customer(
                r["name"],
                r["email"],
                r["phone"],
                r["city"],
                customer_id=r["customer_id"],
            )
            for r in rows
        ]

    def search_customers(self, query: str) -> List[Customer]:
        """Поиск клиентов по имени, email, телефону или городу.

        Parameters
        ----------
        query : str
            Строка поиска.

        Returns
        -------
        list[Customer]
            Список найденных клиентов.
        """
        query_lower = query.strip().lower()
        if not query_lower:
            return self.get_all_customers()
        result = []
        for row in self._read_all("customers"):
            searchable = " ".join(
                str(v) for v in row.values() if v is not None
            ).lower()
            if query_lower in searchable:
                result.append(
                    Customer(
                        row["name"],
                        row["email"],
                        row["phone"],
                        row["city"],
                        customer_id=row["customer_id"],
                    )
                )
        return result

    # --- Заказы ---

    def add_order(self, order: Order) -> int:
        """Добавить заказ в базу.

        Parameters
        ----------
        order : Order
            Заказ для добавления.

        Returns
        -------
        int
            Идентификатор добавленного заказа.
        """
        order_rows = self._read_all("orders")
        item_rows = self._read_all("order_items")
        if order.order_id is None or any(
            r["order_id"] == order.order_id for r in order_rows
        ):
            order.order_id = self._next_id("orders")
        else:
            order_rows = [r for r in order_rows if r["order_id"] != order.order_id]
            item_rows = [r for r in item_rows if r["order_id"] != order.order_id]
        order_rows.append(
            {
                "order_id": order.order_id,
                "customer_id": order.customer.customer_id,
                "order_date": order.order_date,
                "status": order.status,
            }
        )
        next_item_id = self._next_id("order_items")
        for item in order.items:
            item_rows.append(
                {
                    "item_id": next_item_id,
                    "order_id": order.order_id,
                    "product_id": item.product.product_id,
                    "quantity": item.quantity,
                }
            )
            next_item_id += 1
        self._rewrite("orders", order_rows)
        self._rewrite("order_items", item_rows)
        return order.order_id

    def delete_order(self, order_id: int) -> None:
        """Удалить заказ по идентификатору.

        Parameters
        ----------
        order_id : int
            Идентификатор заказа.
        """
        order_rows = [
            r for r in self._read_all("orders") if r["order_id"] != order_id
        ]
        item_rows = [
            r for r in self._read_all("order_items") if r["order_id"] != order_id
        ]
        self._rewrite("orders", order_rows)
        self._rewrite("order_items", item_rows)

    def get_all_orders(self) -> List[Order]:
        """Получить список всех заказов.

        Returns
        -------
        list[Order]
            Список заказов.
        """
        products = {p.product_id: p for p in self.get_all_products()}
        customers = {c.customer_id: c for c in self.get_all_customers()}

        order_rows = self._read_all("orders")
        item_rows = self._read_all("order_items")

        items_by_order: dict = {}
        for row in item_rows:
            items_by_order.setdefault(row["order_id"], []).append(
                (row["product_id"], row["quantity"])
            )

        orders = []
        for row in order_rows:
            customer = customers.get(row["customer_id"])
            if customer is None:
                continue
            order = Order(
                customer=customer,
                order_date=row["order_date"],
                status=row["status"],
                order_id=row["order_id"],
            )
            for product_id, quantity in items_by_order.get(row["order_id"], []):
                product = products.get(product_id)
                if product is not None:
                    order.add_item(product, quantity)
            orders.append(order)
        return orders

    # --- Импорт / экспорт ---

    def export_json(self, path: str) -> None:
        """Экспортировать все данные в файл JSON.

        Parameters
        ----------
        path : str
            Путь к файлу JSON.
        """
        data = {
            "products": [p.to_dict() for p in self.get_all_products()],
            "customers": [c.to_dict() for c in self.get_all_customers()],
            "orders": [o.to_dict() for o in self.get_all_orders()],
        }
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

    def import_json(self, path: str) -> None:
        """Импортировать данные из файла JSON.

        Parameters
        ----------
        path : str
            Путь к файлу JSON.
        """
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        for p_data in data.get("products", []):
            self.add_product(Product.from_dict(p_data))
        for c_data in data.get("customers", []):
            self.add_customer(Customer.from_dict(c_data))
        products = self.get_all_products()
        customers = self.get_all_customers()
        customer_map = {c.customer_id: c for c in customers}
        for o_data in data.get("orders", []):
            customer = customer_map.get(o_data.get("customer_id"))
            if customer is None:
                continue
            order = Order.from_dict(o_data, customer, products)
            self.add_order(order)

    def export_orders_csv(self, path: str) -> None:
        """Экспортировать заказы в файл CSV.

        Parameters
        ----------
        path : str
            Путь к файлу CSV.
        """
        with open(path, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(
                ["order_id", "date", "customer_name", "customer_city", "status", "total"]
            )
            for order in self.get_all_orders():
                writer.writerow(
                    [
                        order.order_id,
                        order.order_date,
                        order.customer.name,
                        order.customer.city,
                        order.status,
                        f"{order.total():.2f}",
                    ]
                )

    def export_customers_csv(self, path: str) -> None:
        """Экспортировать клиентов в файл CSV.

        Parameters
        ----------
        path : str
            Путь к файлу CSV.
        """
        with open(path, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(["customer_id", "name", "email", "phone", "city"])
            for customer in self.get_all_customers():
                writer.writerow(
                    [
                        customer.customer_id,
                        customer.name,
                        customer.email,
                        customer.phone,
                        customer.city,
                    ]
                )

    def export_products_csv(self, path: str) -> None:
        """Экспортировать товары в файл CSV.

        Parameters
        ----------
        path : str
            Путь к файлу CSV.
        """
        with open(path, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(["product_id", "name", "price", "category"])
            for product in self.get_all_products():
                writer.writerow(
                    [product.product_id, product.name, product.price, product.category]
                )

    def import_csv(self, path: str) -> str:
        """Импортировать данные из CSV-файла.

        Формат определяется по заголовку первой строки.

        Parameters
        ----------
        path : str
            Путь к файлу CSV.

        Returns
        -------
        str
            Сообщение о результате импорта.
        """
        with open(path, "r", encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f, delimiter=";")
            rows = list(reader)
        if not rows:
            return "Файл CSV пуст"
        headers = set(rows[0].keys())
        count = 0
        if "product_id" in headers and "category" in headers:
            for row in rows:
                self.add_product(
                    Product(row["name"], float(row["price"]), row["category"])
                )
                count += 1
            return f"Импортировано товаров: {count}"
        if "customer_id" in headers and "city" in headers:
            for row in rows:
                self.add_customer(
                    Customer(row["name"], row["email"], row["phone"], row["city"])
                )
                count += 1
            return f"Импортировано клиентов: {count}"
        return f"Неподдерживаемый формат CSV: {headers}"